"""读取销售表，预览四条但导入全部；问题行保留，不静默删除或补零。"""

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from .common import json_text, js_string, trim_text, utf16_length

ALIASES = {"date": ["销售日期", "日期", "sold_at", "date"],
           "region": ["区域", "地区", "region"], "product": ["商品名称", "商品", "product"],
           "paid": ["实付金额（元）", "实付金额", "paid_amount"],
           "refund": ["退款金额（元）", "退款金额", "refund_amount"]}


def check_size(cells):
    if not cells or len(cells) > 10001 or any(len(row) > 50 for row in cells):
        raise ValueError("空表或超过 10000 行 / 50 列")


def parse_file(data, filename):
    """严格 UTF-8 CSV 或 XLSX；不采用公式缓存值，不接受合并表头和合并数据。"""
    if not data or len(data) > 5 * 1024 * 1024:
        raise ValueError("文件须为非空且不超过 5 MiB")
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        text = data.decode("utf-8-sig", errors="strict")
        cells = [row for row in csv.reader(io.StringIO(text, newline=""), strict=True) if row]
        if cells and any(len(row) != len(cells[0]) for row in cells):
            raise ValueError("CSV 各行列数不一致，请核对分隔符与引号")
        if any(sum(utf16_length(value) for value in row) > 100000 for row in cells):
            raise ValueError("CSV 单条记录过长")
        check_size(cells)
        return [{"name": "CSV", "cells": cells}]
    if extension != ".xlsx" or not data.startswith(b"PK"):
        raise ValueError("只支持 XLSX 与 UTF-8 CSV")
    workbook = load_workbook(io.BytesIO(data), data_only=False)
    try:
        result = []
        for sheet in workbook.worksheets:
            if sheet.max_row > 10001 or sheet.max_column > 50:
                raise ValueError("每张表最多 10000 行数据、50 列")
            if sheet.merged_cells.ranges:
                raise ValueError("请先取消表头和数据区域的合并单元格")
            cells = []
            for row in sheet.iter_rows():
                # 保留公式标记让质量检查明确报错，不把缓存结果当成可信金额。
                cells.append([{"formula": cell.value[1:]} if cell.data_type == "f" else
                              {"error": cell.value} if cell.data_type == "e" else cell.value for cell in row])
            # ExcelJS 空工作表没有行；openpyxl 的默认 A1 空格须规范化为空表。
            if not any(any(value is not None for value in row) for row in cells):
                cells = []
            check_size(cells)
            result.append({"name": sheet.title, "cells": cells})
        return result
    finally:
        workbook.close()


def nonempty(row):
    return any(value is not None and value != "" for value in row)


def preview_sheet(sheet):
    headers = [trim_text(value) if isinstance(value, str) else "" for value in sheet["cells"][0]]
    if any(not header for header in headers) or len(set(headers)) != len(headers):
        raise ValueError(f"{sheet['name']} 的首行须为不重复的文字表头")
    suggested = {key: next((header for header in headers if header in names), "") for key, names in ALIASES.items()}
    data_rows = [row for row in sheet["cells"][1:] if nonempty(row)]
    def preview(value):
        return value.isoformat()[:10] if isinstance(value, (date, datetime)) else value
    return {"name": sheet["name"], "headers": headers, "rowCount": len(data_rows),
            "suggested": suggested, "rows": [[preview(value) for value in row] for row in data_rows[:4]]}


def normalize(sheet, mapping):
    headers = preview_sheet(sheet)["headers"]
    if (set(mapping) != set(ALIASES) or any(value not in headers for value in mapping.values())
            or len(set(mapping.values())) != 5):
        raise ValueError("请选择五个不同且真实存在的字段；金额单位须为元")
    rows, issues, seen = [], [], {}
    for index, cells in enumerate(sheet["cells"][1:], 2):
        if not nonempty(cells):
            continue
        result = {"source_row": index}
        for field, header in mapping.items():
            position = headers.index(header)
            value = cells[position] if position < len(cells) else None
            try:
                if value is None or value == "":
                    raise ValueError("必填值为空")
                if isinstance(value, (dict, list, tuple)):
                    raise ValueError("含公式或复杂单元格，请先核实并导出数值")
                if field == "date":
                    text = value.isoformat()[:10] if isinstance(value, (date, datetime)) else trim_text(js_string(value))
                    try:
                        if not re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", text):
                            raise ValueError()
                        date.fromisoformat(text)
                    except ValueError:
                        raise ValueError("日期须为有效 YYYY-MM-DD") from None
                    result["sold_at"] = text
                elif field in ("paid", "refund"):
                    text = trim_text(js_string(value)).replace(",", "")
                    if not re.fullmatch(r"(0|[1-9][0-9]*)(\.[0-9]{1,2})?", text) or Decimal(text) > 1000000000:
                        raise ValueError("金额须为非负数，最多两位小数，单位元")
                    result["paid_amount" if field == "paid" else "refund_amount"] = f"{Decimal(text):.2f}"
                else:
                    result["product_name" if field == "product" else "region"] = trim_text(js_string(value))
            except ValueError as error:
                issues.append({"row": index, "field": header, "message": str(error)})
        if (result.get("refund_amount") is not None and result.get("paid_amount") is not None
                and Decimal(result["refund_amount"]) > Decimal(result["paid_amount"])):
            issues.append({"row": index, "field": mapping["refund"], "message": "退款金额超过实付金额"})
        key = json_text(cells)
        if key in seen:
            issues.append({"row": index, "field": "整行", "message": f"与第 {seen[key]} 行完全相同，请核对是否重复导入；未自动删除"})
        else:
            seen[key] = index
        rows.append(result)
    if not rows:
        raise ValueError("所选工作表没有明细")
    return {"rows": rows, "issues": issues}
